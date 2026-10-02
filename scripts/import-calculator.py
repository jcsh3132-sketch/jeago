import os
import re
import threading
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
PRICE_FILE_PATTERN = "*판매가격표*.xlsx"
PRINTER_SHEET_PREFIX = "프린터 가격"
CONSUMABLE_SHEET_KEYWORD = "소모품"
RENT_SHEET_NAME = "복합기 임대"
FILTER_SHEET_KEYWORD = "필터"
COPIER_SHEET_KEYWORD = "복사기"
TONER_SHEET_PREFIX = "토너 가격표"
CONSUMABLE_DISCOUNT_RATE = 0.88




def cell(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else ""


def text_value(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def number_value(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return None


# =========================================================
# 가격표 엑셀 파일 읽기
# ---------------------------------------------------------
# 프로젝트 폴더에서 이름에 "판매가격표"가 들어간 .xlsx 중
# 파일명이 가장 뒤인 것(예: 202609_판매가격표.xlsx)을 사용합니다.
# 새 달 가격표를 폴더에 넣으면 자동으로 그 파일을 읽습니다.
# =========================================================
def get_price_file() -> Path:
    configured = os.environ.get("PRICE_FILE")
    if configured:
        path = Path(configured).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"가격표 파일을 찾을 수 없습니다: {path}")
        return path

    candidates = sorted(
        path for path in BASE_DIR.glob(PRICE_FILE_PATTERN) if not path.name.startswith("~$")
    )
    if not candidates:
        raise FileNotFoundError(f"프로젝트 폴더에 가격표 파일({PRICE_FILE_PATTERN})이 없습니다.")
    return candidates[-1]


# =========================================================
# 가격 인상 파일 (예: 프린트, 소모품(2026년 9월 가격표).xlsx)
# ---------------------------------------------------------
# 기본 가격표에는 임대·필터·토너 등 여러 시트가 있고, 인상 파일에는 프린터·소모품만 있습니다.
# 폴더에 "YYYY년 M월 가격표"가 들어간 파일이 있으면 가장 최근 달 파일의
# 최신 월 반출가/출고가로 기본 가격표의 프린터·소모품 가격을 덮어씁니다.
# =========================================================
PRICE_UPDATE_NAME = re.compile(r"(\d{4})년\s*(\d{1,2})월\s*가격표")
MONTH_HEADER = re.compile(r"\d+년\s*\d+월")


def get_price_update_file() -> Path | None:
    candidates = []
    for path in BASE_DIR.glob("*.xlsx"):
        match = PRICE_UPDATE_NAME.search(path.stem)
        if match and not path.name.startswith("~$"):
            candidates.append(((int(match.group(1)), int(match.group(2))), path))
    return max(candidates)[1] if candidates else None


def read_latest_prices(worksheet) -> dict[str, tuple[int | None, int | None]]:
    """'모델명' 머리글 행을 찾아, 가장 오른쪽 'N년 M월' 칸(반출가)과 그 옆 칸(출고가)을 읽습니다."""
    prices: dict[str, tuple[int | None, int | None]] = {}
    model_col = price_col = None
    for row in worksheet.iter_rows(values_only=True):
        texts = [text_value(value) for value in row]
        if "모델명" in texts:
            model_col = texts.index("모델명")
            month_cols = [index for index, text in enumerate(texts) if MONTH_HEADER.search(text)]
            price_col = month_cols[-1] if month_cols else None
            continue
        if model_col is None or price_col is None:
            continue
        model = text_value(cell(row, model_col))
        release_price, sale_price = number_value(cell(row, price_col)), number_value(cell(row, price_col + 1))
        if model and sale_price is not None:
            prices.setdefault(model.upper(), (release_price, sale_price))
    return prices


def read_price_update(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        printer_title = next((name for name in workbook.sheetnames if "프린터" in name), None)
        consumable_title = next((name for name in workbook.sheetnames if CONSUMABLE_SHEET_KEYWORD in name), None)
        match = PRICE_UPDATE_NAME.search(path.stem)
        return {
            "file": path.name,
            "label": f"{match.group(1)}년 {match.group(2)}월",
            "printers": read_latest_prices(workbook[printer_title]) if printer_title else {},
            "consumables": read_latest_prices(workbook[consumable_title]) if consumable_title else {},
        }
    finally:
        workbook.close()


def set_cell(row: list[Any], index: int, value: Any) -> None:
    row.extend([None] * (index + 1 - len(row)))
    row[index] = value


def apply_price_update(
    printer_rows: list[list[Any]], consumable_rows: list[list[Any]], update: dict[str, Any]
) -> dict[str, int]:
    """기본 가격표 행의 반출가/출고가를 인상 파일 값으로 바꿉니다. 프린터 구매가는 출고가 × 0.88 / 0.8로 다시 계산합니다."""
    changed = {"printers": 0, "consumables": 0}
    update["consumableDeltas"] = {}

    for row in printer_rows:
        new = update["printers"].get(text_value(cell(row, 2)).upper())
        if not new or (number_value(cell(row, 6)), number_value(cell(row, 7))) == new:
            continue
        set_cell(row, 6, new[0])
        set_cell(row, 7, new[1])
        if number_value(cell(row, 8)) is not None:
            set_cell(row, 8, new[1] * 0.88)
        if number_value(cell(row, 9)) is not None:
            set_cell(row, 9, new[1] * 0.8)
        changed["printers"] += 1

    for row in consumable_rows:
        model = text_value(cell(row, 3)).upper()
        new = update["consumables"].get(model)
        if not new or (number_value(cell(row, 7)), number_value(cell(row, 8))) == new:
            continue
        update["consumableDeltas"].setdefault(model, new[1] - number_value(cell(row, 8)))
        set_cell(row, 7, new[0])
        set_cell(row, 8, new[1])
        changed["consumables"] += 1

    return changed


_workbook_cache: dict[str, Any] = {}
_workbook_lock = threading.Lock()


def get_workbook_data() -> dict[str, Any]:
    """엑셀 파일은 크기가 커서 수정 시각이 바뀔 때만 다시 읽습니다. 동시 요청이 와도 한 번만 읽습니다."""
    with _workbook_lock:
        return _load_workbook_data()


def _load_workbook_data() -> dict[str, Any]:
    price_file = get_price_file()
    update_file = get_price_update_file()
    cache_key = f"{price_file}:{price_file.stat().st_mtime_ns}"
    if update_file:
        cache_key += f"|{update_file}:{update_file.stat().st_mtime_ns}"
    if _workbook_cache.get("key") == cache_key:
        return _workbook_cache["data"]

    workbook = load_workbook(price_file, read_only=True, data_only=True)
    try:
        def read_sheet(title: str, max_col: int) -> list[list[Any]]:
            return [list(row) for row in workbook[title].iter_rows(max_col=max_col, values_only=True)]

        def find_sheet(match) -> str:
            title = next((name for name in workbook.sheetnames if match(name)), None)
            if title is None:
                raise ValueError(f"{price_file.name}에서 필요한 시트를 찾을 수 없습니다.")
            return title

        printer_title = find_sheet(lambda name: name.startswith(PRINTER_SHEET_PREFIX))
        consumable_title = find_sheet(lambda name: CONSUMABLE_SHEET_KEYWORD in name)
        rent_title = find_sheet(lambda name: name.strip() == RENT_SHEET_NAME)
        filter_title = find_sheet(lambda name: FILTER_SHEET_KEYWORD in name)
        copier_title = find_sheet(lambda name: COPIER_SHEET_KEYWORD in name)
        toner_title = find_sheet(lambda name: name.startswith(TONER_SHEET_PREFIX))

        printer_rows = read_sheet(printer_title, 22)
        consumable_rows, highlighted_consumable_rows = read_consumable_sheet(workbook[consumable_title])
        toner_items = read_toner_sheet(workbook[toner_title])
        rent_rows = read_sheet(rent_title, 7)
        filter_rows = read_sheet(filter_title, 5)
        copier_rows = read_sheet(copier_title, 3)
    finally:
        workbook.close()

    price_update = None
    if update_file:
        price_update = read_price_update(update_file)
        price_update["changed"] = apply_price_update(printer_rows, consumable_rows, price_update)

    printers = merge_copier_models(parse_products(printer_rows), parse_copier_costs(copier_rows))
    resolve_toner_names(toner_items, consumable_rows)
    consumables, unmatched_printers = parse_consumables(
        consumable_rows,
        printers,
        highlighted_consumable_rows,
        forced_models={item["resolved"] for item in toner_items},
    )
    consumables = apply_toner_prices(
        consumables, toner_items, toner_title, price_update["consumableDeltas"] if price_update else None
    )
    data = {
        "file": price_file.name,
        "priceUpdate": (
            {key: price_update[key] for key in ("file", "label", "changed")} if price_update else None
        ),
        "printerSheet": printer_title,
        "consumableSheet": consumable_title,
        "rentSheet": rent_title,
        "filterSheet": filter_title,
        "printers": printers,
        "consumables": consumables,
        "filters": parse_filters(filter_rows),
        "printersWithoutConsumables": unmatched_printers,
        "rentProducts": parse_rent_products(rent_rows),
    }
    _workbook_cache.update(key=cache_key, data=data)
    return data


# =========================================================
# 프린터 가격 시트
# ---------------------------------------------------------
# 원가: 구매가 12% 또는 20% (엑셀 기준 VAT 포함, 출고가 × 0.88 / 0.8)
# 렌탈전용모델은 구매가가 없으므로 반출가(VAT별도)를 원가로 씁니다.
# =========================================================
def parse_products(rows: list[list[Any]]) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    current_category = ""

    for row_number, row in enumerate(rows, start=1):
        category = text_value(cell(row, 1))
        if category:
            current_category = category

        model = text_value(cell(row, 2))
        feature = text_value(cell(row, 3))
        product_type = text_value(cell(row, 4))
        sales_route = text_value(cell(row, 5))
        release_price_ex_vat = number_value(cell(row, 6))
        sale_price_with_vat = number_value(cell(row, 7))
        purchase_12 = number_value(cell(row, 8))
        purchase_20 = number_value(cell(row, 9))
        series = text_value(cell(row, 21))

        if not model or not any((release_price_ex_vat, sale_price_with_vat, purchase_12, purchase_20)):
            continue

        is_rental = "렌탈전용모델" in current_category or sales_route == "렌탈"
        purchase_values = [value for value in (purchase_12, purchase_20) if value is not None]

        if is_rental:
            product_cost = release_price_ex_vat
            cost_source = "반출가(VAT별도)"
            status = "ok" if product_cost is not None else "missing_cost"
        elif len(purchase_values) > 1:
            product_cost = None
            cost_source = "구매가 중복"
            status = "duplicate_cost"
        elif len(purchase_values) == 1:
            product_cost = purchase_values[0]
            cost_source = "구매가 12%(VAT포함)" if purchase_12 is not None else "구매가 20%(VAT포함)"
            status = "ok"
        else:
            product_cost = None
            cost_source = "구매가 없음"
            status = "missing_cost"

        products.append(
            {
                "kind": "printer",
                "model": model,
                "category": current_category,
                "feature": feature,
                "productType": product_type,
                "salesRoute": sales_route,
                "series": series,
                "releasePriceExVat": release_price_ex_vat,
                "salePriceWithVat": sale_price_with_vat,
                "isRental": is_rental,
                "productCost": product_cost,
                "costSource": cost_source,
                "status": status,
                "rowNumber": row_number,
                "purchase12": purchase_12,
                "purchase20": purchase_20,
            }
        )

    return products


# =========================================================
# 복사기가격 시트 (대형 복합기)
# ---------------------------------------------------------
# 스토어별 구역마다 같은 모델이 반복되므로, 모델별로 처음 나오는 매입금액(C열)을 씁니다.
# - 프린터 시트에 있는데 원가가 비어 있는 복합기는 매입금액으로 채웁니다.
# - 프린터 시트에 없는 복합기(SL-X3220NR 등)는 새 모델로 추가합니다.
# 매입금액이 비어 있으면 원가 없음으로 표시합니다.
# =========================================================
def parse_copier_costs(rows: list[list[Any]]) -> dict[str, dict[str, Any]]:
    copiers: dict[str, dict[str, Any]] = {}

    for row_number, row in enumerate(rows, start=1):
        model = text_value(cell(row, 1))
        if "-" not in model:
            continue

        cost = number_value(cell(row, 2))
        existing = copiers.get(model.upper())
        if existing is None:
            copiers[model.upper()] = {"model": model, "cost": cost, "rowNumber": row_number}
        elif existing["cost"] is None and cost is not None:
            existing.update(cost=cost, rowNumber=row_number)

    return copiers


def copier_category(model: str) -> str:
    if model.upper().startswith("SL-X"):
        return "A3 컬러 DMFP"
    if model.upper().startswith("SL-K"):
        return "A3 흑백 DMFP"
    return "복합기"


def merge_copier_models(
    printers: list[dict[str, Any]], copiers: dict[str, dict[str, Any]]
) -> list[dict[str, Any]]:
    registered = {printer["model"].upper() for printer in printers}

    for printer in printers:
        copier = copiers.get(printer["model"].upper())
        if printer["status"] == "missing_cost" and not printer["isRental"] and copier and copier["cost"] is not None:
            printer.update(productCost=copier["cost"], costSource="매입금액(복사기가격)", status="ok")

    extras = [
        {
            "kind": "printer",
            "model": copier["model"],
            "category": copier_category(copier["model"]),
            "feature": "",
            "productType": "복사기가격 시트",
            "salesRoute": "",
            "series": "",
            "releasePriceExVat": None,
            "salePriceWithVat": None,
            "isRental": False,
            "productCost": copier["cost"],
            "costSource": "매입금액(복사기가격)" if copier["cost"] is not None else "매입금액 없음",
            "status": "ok" if copier["cost"] is not None else "missing_cost",
            "rowNumber": copier["rowNumber"],
            "sourceSheet": "복사기가격",
            "purchase12": None,
            "purchase20": None,
        }
        for key, copier in copiers.items()
        if key not in registered
    ]
    return printers + extras


# =========================================================
# 소모품 시트
# ---------------------------------------------------------
# "대상 모델" 칸은 그룹의 첫 행에만 적혀 있으므로 아래 행에 이어서 적용합니다.
# 대상 모델 중 프린터 가격 시트에 등록된 프린터가 하나라도 있는 소모품만 남깁니다.
# 원가: 출고가(VAT포함) × 0.88 (시트의 "할인 12%(VAT포함)" 기준)
# =========================================================
HIGHLIGHT_COLOR = "FFFFFF00"  # 노란색: 우리가 취급(매입)하는 소모품
RENTAL_CONSUMABLE_GROUP = "렌탈전용모델"


def is_highlighted_cell(target) -> bool:
    fill = getattr(target, "fill", None)
    color = getattr(fill, "fgColor", None)
    return color is not None and color.type == "rgb" and color.rgb == HIGHLIGHT_COLOR


def read_consumable_sheet(worksheet) -> tuple[list[list[Any]], set[int]]:
    """소모품 시트 값과 함께, 모델명(D열) 칸이 노란색인 행 번호를 돌려줍니다."""
    rows: list[list[Any]] = []
    highlighted: set[int] = set()
    for row_number, cells in enumerate(worksheet.iter_rows(max_col=12), start=1):
        rows.append([c.value for c in cells])
        if len(cells) > 3 and is_highlighted_cell(cells[3]):
            highlighted.add(row_number)
    return rows, highlighted


MODEL_TOKEN = re.compile(r"^(?:([A-Z]+)-)?([A-Z]*)(\d+[A-Z]*(?:~[A-Z]+)?)$")
SERIES_TOKEN = re.compile(r"^([A-Z]\d{2}R)(?:시리즈|SERIES)?$")
MODEL_SUFFIXES = "KRM|GOV|HYP|DCS|EXP"


def parse_target_models(text: str) -> tuple[list[str], list[str]]:
    """'SL-C435, C436\\nSL-M2630/M2680 시리즈' 같은 문자열을 모델명과 시리즈 코드로 나눕니다."""
    cleaned = re.sub(r"\([^)]*\)", " ", text.upper())
    # SL-M3830ND/KRM 같은 모델 접미사는 구분자 "/"와 헷갈리지 않도록 잠시 "~"로 바꿉니다.
    cleaned = re.sub(rf"/(?=(?:{MODEL_SUFFIXES})\b)", "~", cleaned)
    models: list[str] = []
    series: list[str] = []
    prefix, letters = "", ""

    for raw in re.split(r"[,\n/.\s]+", cleaned):
        token = raw.strip()
        if not token:
            continue
        series_match = SERIES_TOKEN.match(token)
        if series_match:
            series.append(series_match.group(1))
            continue

        token = re.sub(r"(시리즈|SERIES)$", "", token)
        match = MODEL_TOKEN.match(token)
        if not match:
            continue
        token_prefix, token_letters, number = match.groups()
        if token_prefix:
            # SL 모델은 항상 영문 계열(SL-M, SL-C...)이 붙으므로 "SL-9500LX" 같은 누락은 앞 모델 계열을 이어 씁니다.
            if token_letters or token_prefix != "SL" or prefix != "SL":
                letters = token_letters
            prefix = token_prefix
        elif token_letters:
            letters = token_letters
            prefix = prefix or "SL"
        if not prefix:
            continue
        models.append(f"{prefix}-{letters}{number}".replace("~", "/"))

    return models, series


def printer_matches(printer_model: str, target: str) -> bool:
    """대상 모델명이 등록 프린터와 같은지 봅니다. 숫자로 끝나는 대상은 뒤 옵션 문자(ND, W 등)까지 허용합니다."""
    model = printer_model.upper()
    if model == target:
        return True
    if not model.startswith(target):
        return False
    rest = model[len(target):]
    if rest.startswith("/"):
        return True
    return target[-1].isdigit() and not rest[0].isdigit()


def parse_consumables(
    rows: list[list[Any]],
    printers: list[dict[str, Any]],
    highlighted_rows: set[int],
    forced_models: set[str] = frozenset(),
) -> tuple[list[dict[str, Any]], list[str]]:
    """렌탈전용모델 구역은 대상 프린터 대조 대신 노란색으로 표시된 소모품만 등록합니다.
    forced_models(토너 가격표에 있는 모델)는 대조 결과와 관계없이 등록합니다."""
    consumables: dict[str, dict[str, Any]] = {}
    matched_printers: set[str] = set()
    group, subgroup, target_text, compatible = "", "", "", []

    for row_number, row in enumerate(rows, start=1):
        if text_value(cell(row, 1)):
            group = text_value(cell(row, 1))
        if text_value(cell(row, 2)):
            subgroup = text_value(cell(row, 2))

        model = text_value(cell(row, 3))
        release_price_ex_vat = number_value(cell(row, 7))
        sale_price_with_vat = number_value(cell(row, 8))
        if not model or sale_price_with_vat is None:
            continue

        if text_value(cell(row, 10)):
            target_text = text_value(cell(row, 10))
            target_models, target_series = parse_target_models(target_text)
            compatible = [
                printer["model"]
                for printer in printers
                if printer["series"] in target_series
                or any(printer_matches(printer["model"], target) for target in target_models)
            ]

        is_rental_group = subgroup == RENTAL_CONSUMABLE_GROUP
        is_highlighted = row_number in highlighted_rows
        if model.upper() in forced_models:
            pass
        elif is_rental_group:
            if not is_highlighted:
                continue
        elif not compatible:
            continue
        matched_printers.update(compatible)

        existing = consumables.get(model.upper())
        if existing:
            existing["compatibleModels"] += [m for m in compatible if m not in existing["compatibleModels"]]
            continue

        consumables[model.upper()] = {
            "kind": "consumable",
            "model": model,
            "category": f"{group} {subgroup}".strip(),
            "status": "ok",
            "supplyStatus": text_value(cell(row, 4)),
            "color": text_value(cell(row, 5)),
            "yield": text_value(cell(row, 6)),
            "releasePriceExVat": release_price_ex_vat,
            "salePriceWithVat": sale_price_with_vat,
            "productCost": round(sale_price_with_vat * CONSUMABLE_DISCOUNT_RATE),
            "costSource": "구매가 12%(출고가×0.88, VAT포함)",
            "compatibleModels": list(compatible),
            "targetText": target_text,
            "note": text_value(cell(row, 11)),
            "isRental": is_rental_group,
            "isHighlighted": is_highlighted,
            "rowNumber": row_number,
        }

    unmatched = [
        printer["model"]
        for printer in printers
        if printer["model"] not in matched_printers and printer["productType"] != "옵션"
    ]
    return list(consumables.values()), unmatched


# =========================================================
# 토너 가격표 시트
# ---------------------------------------------------------
# G열 "출고가 = 공급가(부가세 포함)" 중 노란색 칸이 실제로 우리가 받는 가격입니다.
# - 해당 소모품의 원가를 이 공급가로 바꿉니다.
# - "4색 세트" 행은 바로 위 모델들을 묶은 세트 상품으로 등록합니다.
# - 소모품 시트와 이름이 조금 다르면(CLT-K506/RTL ↔ CLT-K506L/RTL) 소모품 시트 이름으로 맞춥니다.
# =========================================================
def read_toner_sheet(worksheet) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    group = ""
    for row_number, cells in enumerate(worksheet.iter_rows(max_col=10), start=1):
        values = [c.value for c in cells]
        if text_value(cell(values, 1)):
            group = text_value(cell(values, 1))
        model = text_value(cell(values, 2))
        supply = number_value(cell(values, 6))
        if not model or supply is None or len(cells) <= 6 or not is_highlighted_cell(cells[6]):
            continue
        items.append(
            {
                "model": model,
                "group": group,
                "yield": text_value(cell(values, 4)),
                "supplyPrice": supply,
                "rowNumber": row_number,
            }
        )
    return items


def resolve_toner_names(toner_items: list[dict[str, Any]], consumable_rows: list[list[Any]]) -> None:
    known = {text_value(cell(row, 3)).upper() for row in consumable_rows if text_value(cell(row, 3))}
    for item in toner_items:
        name = item["model"].upper()
        with_l = re.sub(r"(\d)/", r"\1L/", name)
        item["resolved"] = name if name in known or with_l not in known else with_l


def toner_cost_source(delta: int) -> str:
    if not delta:
        return "토너 가격표 공급가(VAT포함)"
    return f"토너 가격표 공급가(VAT포함) {delta:+,}원 인상분 반영"


def apply_toner_prices(
    consumables: list[dict[str, Any]],
    toner_items: list[dict[str, Any]],
    toner_title: str,
    deltas: dict[str, int] | None = None,
) -> list[dict[str, Any]]:
    """deltas: 가격 인상 파일의 출고가 증감액. 공급가에 그 차액만 더합니다(세트는 구성품 차액 합계)."""
    deltas = deltas or {}
    by_model = {item["model"].upper(): item for item in consumables}
    sets: list[dict[str, Any]] = []
    members: list[dict[str, Any]] = []
    previous_row = 0

    for item in toner_items:
        if item["rowNumber"] != previous_row + 1:
            members = []
        previous_row = item["rowNumber"]

        if "세트" in item["model"]:
            if not members:
                continue
            base = re.sub(r"^(CLT-)[KCMY]", r"\1", members[0]["model"])
            compatible: list[str] = []
            for member in members:
                compatible += [m for m in member.get("compatibleModels", []) if m not in compatible]
            sets.append(
                {
                    "kind": "consumable",
                    "model": f"{base} {item['model']}",
                    "category": members[0].get("category", item["group"]),
                    "status": "ok",
                    "supplyStatus": members[0].get("supplyStatus", ""),
                    "color": "4색 세트",
                    "yield": "",
                    "releasePriceExVat": None,
                    "salePriceWithVat": None,
                    "productCost": item["supplyPrice"] + sum(member.get("priceDelta", 0) for member in members),
                    "costSource": toner_cost_source(sum(member.get("priceDelta", 0) for member in members)),
                    "compatibleModels": compatible,
                    "targetText": members[0].get("targetText", ""),
                    "note": " + ".join(member["model"] for member in members),
                    "isRental": members[0].get("isRental", False),
                    "isHighlighted": True,
                    "sourceSheet": toner_title,
                    "rowNumber": item["rowNumber"],
                }
            )
            members = []
            continue

        consumable = by_model.get(item["resolved"])
        if consumable is None:
            continue
        delta = deltas.get(item["resolved"], 0)
        consumable.update(
            productCost=item["supplyPrice"] + delta,
            costSource=toner_cost_source(delta),
            priceDelta=delta,
            isHighlighted=True,
            sourceSheet=toner_title,
            rowNumber=item["rowNumber"],
        )
        members.append(consumable)

    return consumables + sets


# =========================================================
# 필터판매금액 시트 (공기청정기 필터)
# ---------------------------------------------------------
# 원가: 매입가 (D열). 임시책정금액(E열)은 참고용으로 함께 보냅니다.
# =========================================================
def parse_filters(rows: list[list[Any]]) -> list[dict[str, Any]]:
    filters: list[dict[str, Any]] = []

    for row_number, row in enumerate(rows, start=1):
        model = text_value(cell(row, 1))
        if not model or model == "품명":
            continue

        purchase_price = number_value(cell(row, 3))
        filters.append(
            {
                "kind": "filter",
                "model": model,
                "category": "공청기 필터",
                "stock": text_value(cell(row, 2)),
                "purchasePrice": purchase_price,
                "suggestedPrice": number_value(cell(row, 4)),
                "productCost": purchase_price,
                "costSource": "매입가",
                "status": "ok" if purchase_price is not None else "missing_cost",
                "rowNumber": row_number,
            }
        )

    return filters


# =========================================================
# 복합기 임대 시트
# =========================================================
def parse_rent_products(rows: list[list[Any]]) -> list[dict[str, Any]]:
    products: list[dict[str, Any]] = []
    machine_type = ""

    for row_number, row in enumerate(rows, start=1):
        section = text_value(cell(row, 1))
        if section == "컬러 복합기":
            machine_type = "color"
        elif section == "흑백 복합기":
            machine_type = "mono"

        model = text_value(cell(row, 2))
        speed = text_value(cell(row, 3))
        base_rent = number_value(cell(row, 4))
        if machine_type not in ("color", "mono") or not model or base_rent is None:
            continue

        products.append(
            {
                "model": model,
                "machineType": machine_type,
                "machineTypeName": "컬러 복합기" if machine_type == "color" else "흑백 복합기",
                "speed": speed,
                "baseRent": base_rent,
                "blackRate": 8,
                "colorRate": 70 if machine_type == "color" else 0,
                "deposit": 300000 if machine_type == "color" else 200000,
                "usedDiscount": 10000 if machine_type == "color" else 7000,
                "rowNumber": row_number,
            }
        )

    return products



if __name__ == "__main__":
    import argparse
    import hashlib
    import json
    cli = argparse.ArgumentParser(description="Import best.zip price workbooks without modifying them; requires openpyxl 3.1.5.")
    cli.add_argument("source", type=Path)
    cli.add_argument("output", type=Path)
    args = cli.parse_args()
    BASE_DIR = args.source.resolve()
    data = get_workbook_data()
    products = data["printers"] + data["consumables"] + data["filters"]
    sale = {"file": data["file"], "priceUpdate": data["priceUpdate"], "sheet": data["printerSheet"], "count": len(products), "printerCount": len(data["printers"]), "consumableCount": len(data["consumables"]), "filterCount": len(data["filters"]), "printersWithoutConsumables": data["printersWithoutConsumables"], "products": products}
    rent = {"file": data["file"], "sheet": data["rentSheet"], "count": len(data["rentProducts"]), "products": data["rentProducts"]}
    sources = [{"file": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in [get_price_file(), get_price_update_file()] if p]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps({"sources": sources, "sale": sale, "rent": rent}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"printers": sale["printerCount"], "consumables": sale["consumableCount"], "filters": sale["filterCount"], "rent": rent["count"], "update": data["priceUpdate"]}, ensure_ascii=False))
